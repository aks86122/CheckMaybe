"""Hidden Cost Tracker prototype (dark theme): debts & installments with real-APR back-calc,
subscriptions with free-trial cancel-by dates, and a dashboard of monthly committed money."""
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule, CellIsRule, DataBarRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import PieChart, BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.comments import Comment

BG, PANEL, INPUT, LINE = "15171F", "1E2130", "262A3D", "33384F"
TXT, MUTED, BLUE, ORANGE, RED, AMBER, GREEN, TEAL = "EDEBF5", "9A98AE", "8AB4FF", "F08A4B", "FF5C5C", "F2B544", "5FD38D", "4FC3C7"
F = lambda c=TXT, b=False, s=10, i=False: Font(name="Arial", color=c, bold=b, size=s, italic=i)
fill = lambda c: PatternFill("solid", fgColor=c)
thin = Side(style="thin", color=LINE); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
CUR, PCT, DATE = '$#,##0.00;($#,##0.00);"-"', '0.0%;-0.0%;"-"', 'mmm d, yyyy'

wb = Workbook()
def sheet(name, widths, first=False):
    ws = wb.active if first else wb.create_sheet(name)
    ws.title = name; ws.sheet_view.showGridLines = False
    for col, w in widths.items(): ws.column_dimensions[col].width = w
    for r in range(1, 80):
        for c in range(1, 30): ws.cell(r, c).fill = fill(BG)
    return ws

def title(ws, text, sub):
    ws["B2"] = text; ws["B2"].font = F(TXT, True, 18)
    ws["B3"] = sub; ws["B3"].font = F(MUTED, False, 10, True)
    ws.row_dimensions[2].height = 28

def header(ws, row, col, labels, color=ORANGE):
    for i, l in enumerate(labels):
        c = ws.cell(row, col + i, l); c.font = F("15171F", True, 9); c.fill = fill(color)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = BOX
    ws.row_dimensions[row].height = 32

def cell(ws, ref, v, fmt=None, inp=False, bold=False, color=None, align="left"):
    c = ws[ref]; c.value = v; c.border = BOX
    c.fill = fill(INPUT if inp else PANEL); c.font = F(color or (BLUE if inp else TXT), bold)
    c.alignment = Alignment(horizontal=align, vertical="center")
    if fmt: c.number_format = fmt
    return c

# ---------------- Settings ----------------
st = sheet("Settings", {"A": 3, "B": 34, "C": 16, "D": 60})
title(st, "Settings", "Blue cells are yours to edit.")
rows = [("Monthly take-home income", 4200, CUR, "Your income after tax, per month."),
        ("High-cost APR (red flag)", 0.20, PCT, "Any debt at or above this APR is flagged HIGH."),
        ("Check APR (amber flag)", 0.10, PCT, "At or above this APR is flagged CHECK."),
        ("Warn me this many days ahead", 7, "0", "Trials and bills due within this many days are highlighted."),
        ("Today's date", "=TODAY()", DATE, "Updates automatically. Every date calculation uses it.")]
for i, (lab, val, fmt, note) in enumerate(rows):
    r = 5 + i
    cell(st, f"B{r}", lab); cell(st, f"C{r}", val, fmt, inp=(i < 4), align="right")
    st[f"D{r}"] = note; st[f"D{r}"].font = F(MUTED, False, 9, True)
INCOME, RED_APR, AMB_APR, WARN, TODAY = "Settings!$C$5", "Settings!$C$6", "Settings!$C$7", "Settings!$C$8", "Settings!$C$9"
st["B12"] = "Lists used by the drop-downs"; st["B12"].font = F(MUTED, True, 9)
lists = {"B": ["Credit card", "Loan", "Installment", "BNPL", "Family / friend", "Other"],
         "C": ["Weekly", "Monthly", "Quarterly", "Yearly"],
         "D": ["Streaming", "Software", "Fitness", "Shopping", "News", "Cloud", "Food", "Other"]}
for col, items in lists.items():
    for j, v in enumerate(items): st[f"{col}{13 + j}"] = v; st[f"{col}{13 + j}"].font = F(MUTED, False, 9)

# ---------------- Debts ----------------
de = sheet("Debts & Installments", {"A": 3, "B": 22, "C": 14, "D": 13, "E": 12, "F": 11, "G": 10, "H": 11, "I": 13,
                                    "J": 12, "K": 11, "L": 13, "M": 13, "N": 13, "O": 14, "P": 10, "Q": 10, "R": 11})
title(de, "Debts & Installments", "Enter what you know (blue). The real APR is worked out from your monthly payment when you don't know it.")
cols = ["Name", "Type", "Original amount", "Monthly payment", "Total payments", "Payments made", "Stated APR (if known)",
        "Balance now (if known)", "Real APR", "Payments left", "Balance (est.)", "Still to pay", "Interest still to pay",
        "Debt-free date", "Avalanche order", "Snowball order", "Flag"]
header(de, 5, 2, cols[:8], BLUE); header(de, 5, 10, cols[8:], ORANGE)
sample = [("Credit card", "Credit card", None, 120, None, None, 0.2499, 3200),
          ("Car loan", "Loan", None, 330, None, None, 0.069, 14500),
          ("Phone installment plan", "Installment", 1200, 67, 24, 6, None, None),
          ("Sofa store financing", "Installment", 2000, 105, 24, 3, None, None),
          ("Laptop pay-in-4", "BNPL", 800, 200, 4, 1, None, None),
          ("Loan from family", "Family / friend", None, 100, None, None, 0, 1500)]
N0, N1 = 6, 25
for r in range(N0, N1 + 1):
    s = sample[r - N0] if r - N0 < len(sample) else (None,) * 8
    for j, (col, fmt) in enumerate(zip("BCDEFGHI", [None, None, CUR, CUR, "0", "0", PCT, CUR])):
        cell(de, f"{col}{r}", s[j], fmt, inp=True, align="left" if j < 2 else "right")
    has = f"AND($B{r}<>\"\",$E{r}>0)"
    apr = f'IF(NOT({has}),"",IF(ISNUMBER($H{r}),$H{r},IF(AND($D{r}>0,$F{r}>0),IFERROR(RATE($F{r},-$E{r},$D{r})*12,""),"")))'
    f = {"J": f"={apr}",
         "K": f'=IF(NOT({has}),"",IF($F{r}>0,MAX(0,$F{r}-N($G{r})),IF(AND(ISNUMBER($I{r}),ISNUMBER($J{r})),IF($J{r}=0,ROUNDUP($I{r}/$E{r},0),IFERROR(ROUNDUP(NPER($J{r}/12,-$E{r},$I{r}),0),"Never")),"")))',
         "L": f'=IF(NOT({has}),"",IF(ISNUMBER($I{r}),$I{r},IF(AND(ISNUMBER($K{r}),ISNUMBER($J{r})),IF($J{r}=0,$E{r}*$K{r},PV($J{r}/12,$K{r},-$E{r})),"")))',
         "M": f'=IF(ISNUMBER($K{r}),$E{r}*$K{r},"")',
         "N": f'=IF(AND(ISNUMBER($M{r}),ISNUMBER($L{r})),MAX(0,$M{r}-$L{r}),"")',
         "O": f'=IF(ISNUMBER($K{r}),EDATE({TODAY},$K{r}),"")',
         "P": f'=IF(ISNUMBER($J{r}),COUNTIF($J${N0}:$J${N1},">"&$J{r})+1,"")',
         "Q": f'=IF(ISNUMBER($L{r}),COUNTIF($L${N0}:$L${N1},"<"&$L{r})+1,"")',
         "R": f'=IF(NOT(ISNUMBER($J{r})),"",IF($J{r}>={RED_APR},"HIGH",IF($J{r}>={AMB_APR},"CHECK","OK")))'}
    for col, fx in f.items():
        fmt = {"J": PCT, "K": "0", "L": CUR, "M": CUR, "N": CUR, "O": DATE, "P": "0", "Q": "0"}.get(col)
        cell(de, f"{col}{r}", fx, fmt, align="center" if col in "KPQR" else "right")
de["J5"].comment = Comment("Stated APR if you entered one. Otherwise worked out with RATE() from original amount, monthly payment and number of payments. An estimate: fees charged separately are not included. Your contract is the final word.", "CheckMaybe")
dv = DataValidation(type="list", formula1="=Settings!$B$13:$B$18", allow_blank=True); de.add_data_validation(dv); dv.add(f"C{N0}:C{N1}")
rng = f"R{N0}:R{N1}"
de.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"HIGH"'], fill=fill("5A1E24"), font=F(RED, True)))
de.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"CHECK"'], fill=fill("4A3A15"), font=F(AMBER, True)))
de.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"OK"'], fill=fill("173A28"), font=F(GREEN, True)))
de.conditional_formatting.add(f"J{N0}:J{N1}", FormulaRule(formula=[f'AND(ISNUMBER($J{N0}),$J{N0}>={RED_APR})'], font=F(RED, True)))
de.freeze_panes = "C6"

# ---------------- Subscriptions ----------------
su = sheet("Subscriptions", {"A": 3, "B": 20, "C": 12, "D": 11, "E": 12, "F": 13, "G": 13, "H": 11, "I": 9, "J": 11,
                             "K": 10, "L": 12, "M": 12, "N": 14, "O": 10, "P": 14, "Q": 10, "R": 16, "S": 6})
title(su, "Subscriptions & Free Trials", "Untick Keep? on anything you'd cancel to see what you'd save. Free trials get a cancel-by countdown.")
scol = ["Name", "Category", "Price", "Billing", "Start date", "Paid with", "Auto-renew?", "Keep?", "Free trial?", "Trial days",
        "Monthly cost", "Yearly cost", "Next charge", "Days to charge", "Trial: cancel by", "Trial days left", "Alert"]
header(su, 5, 2, scol[:10], BLUE); header(su, 5, 12, scol[10:], ORANGE)
d = lambda y, m, dd: dt.date(y, m, dd)
ssample = [("Netflix", "Streaming", 15.49, "Monthly", d(2025, 3, 5), "Credit card", "Yes", "Yes", "No", None),
           ("Spotify", "Streaming", 11.99, "Monthly", d(2024, 11, 20), "PayPal", "Yes", "Yes", "No", None),
           ("Gym membership", "Fitness", 45, "Monthly", d(2025, 1, 10), "Debit card", "Yes", "No", "No", None),
           ("Photo editing app", "Software", 59.99, "Monthly", d(2026, 2, 14), "Credit card", "Yes", "Yes", "No", None),
           ("Online shopping club", "Shopping", 139, "Yearly", d(2025, 10, 2), "Credit card", "Yes", "Yes", "No", None),
           ("Cloud storage", "Cloud", 2.99, "Monthly", d(2025, 6, 1), "Apple Pay", "Yes", "Yes", "No", None),
           ("Meal kit", "Food", 69.99, "Weekly", d(2026, 9, 20), "Credit card", "Yes", "Yes", "Yes", 14),
           ("News site", "News", 4.99, "Monthly", d(2026, 9, 15), "PayPal", "Yes", "No", "Yes", 30)]
S0, S1 = 6, 35
for r in range(S0, S1 + 1):
    s = ssample[r - S0] if r - S0 < len(ssample) else (None,) * 10
    for j, (col, fmt) in enumerate(zip("BCDEFGHIJK", [None, None, CUR, None, DATE, None, None, None, None, "0"])):
        cell(su, f"{col}{r}", s[j], fmt, inp=True, align="right" if col in "DFK" else "left")
    live = f'AND($B{r}<>"",$D{r}>0)'
    step = f'IF($E{r}="Monthly",1,IF($E{r}="Quarterly",3,IF($E{r}="Yearly",12,0)))'
    first = f'IF($J{r}="Yes",$F{r}+N($K{r}),$F{r})'                                  # first paid charge
    kk = f'IF({first}>={TODAY},0,CEILING(DATEDIF({first},{TODAY},"M")/{step},1))'
    mnext = f'IF(EDATE({first},{step}*({kk}))>={TODAY},EDATE({first},{step}*({kk})),EDATE({first},{step}*({kk}+1)))'
    wnext = f'IF({first}>={TODAY},{first},{first}+7*CEILING(({TODAY}-{first})/7,1))'
    f = {"L": f'=IF(NOT({live}),"",IF($E{r}="Weekly",$D{r}*52/12,IF($E{r}="Quarterly",$D{r}/3,IF($E{r}="Yearly",$D{r}/12,$D{r}))))',
         "M": f'=IF(ISNUMBER($L{r}),$L{r}*12,"")',
         "N": f'=IF(OR(NOT({live}),$F{r}=""),"",IF($E{r}="Weekly",{wnext},{mnext}))',
         "O": f'=IF(ISNUMBER($N{r}),$N{r}-{TODAY},"")',
         "P": f'=IF(AND({live},$J{r}="Yes",ISNUMBER($K{r})),$F{r}+$K{r}-1,"")',
         "Q": f'=IF(ISNUMBER($P{r}),$P{r}-{TODAY},"")',
         "R": f'=IF(NOT({live}),"",IF(AND(ISNUMBER($Q{r}),$Q{r}<0),"Trial ended",IF(AND(ISNUMBER($Q{r}),$Q{r}<={WARN}),"Cancel soon?",IF($I{r}="No","Cancel",IF(AND(ISNUMBER($O{r}),$O{r}<={WARN}),"Charge soon","OK")))))'}
    for col, fx in f.items():
        fmt = {"L": CUR, "M": CUR, "N": DATE, "O": "0", "P": DATE, "Q": "0"}.get(col)
        cell(su, f"{col}{r}", fx, fmt, align="center" if col in "OQR" else "right")
for rngcol, src in (("C", "=Settings!$D$13:$D$20"), ("E", "=Settings!$C$13:$C$16")):
    v = DataValidation(type="list", formula1=src, allow_blank=True); su.add_data_validation(v); v.add(f"{rngcol}{S0}:{rngcol}{S1}")
yn = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True); su.add_data_validation(yn); yn.add(f"H{S0}:J{S1}")
R = f"R{S0}:R{S1}"
su.conditional_formatting.add(R, CellIsRule(operator="equal", formula=['"Cancel soon?"'], fill=fill("5A1E24"), font=F(RED, True)))
su.conditional_formatting.add(R, CellIsRule(operator="equal", formula=['"Charge soon"'], fill=fill("4A3A15"), font=F(AMBER, True)))
su.conditional_formatting.add(R, CellIsRule(operator="equal", formula=['"Cancel"'], fill=fill("2B2E3A"), font=F(MUTED, True)))
su.conditional_formatting.add(R, CellIsRule(operator="equal", formula=['"OK"'], fill=fill("173A28"), font=F(GREEN, True)))
su.conditional_formatting.add(f"B{S0}:Q{S1}", FormulaRule(formula=[f'$I{S0}="No"'], font=F("6E6C80", False, 10, True)))
su.conditional_formatting.add(f"Q{S0}:Q{S1}", FormulaRule(formula=[f'AND(ISNUMBER($Q{S0}),$Q{S0}<=3)'], fill=fill("5A1E24"), font=F(RED, True)))
su.freeze_panes = "C6"

# ---------------- Dashboard ----------------
db = sheet("Dashboard", {"A": 3, "B": 26, "C": 16, "D": 4, "E": 26, "F": 16, "G": 4, "H": 26, "I": 16}, first=True)
wb.move_sheet("Dashboard", -(len(wb.sheetnames) - 1))
db["B2"] = "HIDDEN COST TRACKER"; db["B2"].font = F(ORANGE, True, 20)
db["B3"] = "See what your debts, installments and subscriptions really cost you."; db["B3"].font = F(MUTED, False, 10, True)
DS, SS = "'Debts & Installments'", "Subscriptions"
dr = lambda c: f"{DS}!${c}${N0}:${c}${N1}"; sr = lambda c: f"{SS}!${c}${S0}:${c}${S1}"
def kpi(r, c, label, fx, fmt, color=TXT):
    lc, vc = db.cell(r, c), db.cell(r + 1, c)
    lc.value = label; lc.font = F(MUTED, True, 9); lc.fill = fill(PANEL)
    vc.value = fx; vc.font = F(color, True, 20); vc.fill = fill(PANEL); vc.number_format = fmt
    db.cell(r, c + 1).fill = fill(PANEL); db.cell(r + 1, c + 1).fill = fill(PANEL)
    db.row_dimensions[r + 1].height = 30
kpi(5, 2, "MONTHLY INCOME", f"={INCOME}", CUR)
kpi(5, 5, "ALREADY COMMITTED EACH MONTH", "=C22", CUR, ORANGE)
kpi(5, 8, "SHARE OF INCOME COMMITTED", "=IF(C18>0,C22/C18,0)", PCT, ORANGE)
kpi(8, 2, "HIGH-COST DEBTS (RED FLAG)", f'=COUNTIF({dr("R")},"HIGH")', "0", RED)
kpi(8, 5, "TRIALS TO CANCEL SOON", f'=COUNTIF({sr("R")},"Cancel soon?")', "0", AMBER)
kpi(8, 8, "YEARLY SAVINGS IF YOU CANCEL", f'=SUMIFS({sr("M")},{sr("I")},"No")', CUR, GREEN)
kpi(11, 2, "HIGHEST REAL APR", f'=IF(COUNT({dr("J")})=0,"",MAX({dr("J")}))', PCT, RED)
kpi(11, 5, "…ON", f'=IF(COUNT({dr("J")})=0,"",INDEX({dr("B")},MATCH(MAX({dr("J")}),{dr("J")},0)))', "@", TXT)
kpi(11, 8, "INTEREST STILL TO PAY (ALL DEBTS)", f'=SUM({dr("N")})', CUR, RED)
db["E12"].font = F(TXT, True, 13)
# where the money goes (chart data)
db["B16"] = "WHERE YOUR MONTHLY INCOME GOES"; db["B16"].font = F(TEAL, True, 10)
parts = [("Loans & cards", f'=SUMIFS({dr("E")},{dr("C")},"Credit card")+SUMIFS({dr("E")},{dr("C")},"Loan")+SUMIFS({dr("E")},{dr("C")},"Family / friend")+SUMIFS({dr("E")},{dr("C")},"Other")'),
         ("Installments & BNPL", f'=SUMIFS({dr("E")},{dr("C")},"Installment")+SUMIFS({dr("E")},{dr("C")},"BNPL")'),
         ("Subscriptions", f'=SUM({sr("L")})')]
cell(db, "B17", "Loans & cards"); cell(db, "C17", parts[0][1], CUR, align="right")
db["B18"], db["C18"] = "Income", f"={INCOME}"
for i, (lab, fx) in enumerate(parts):
    r = 19 + i; cell(db, f"B{r}", lab); cell(db, f"C{r}", fx, CUR, align="right")
db["B17"].value, db["C17"].value = None, None
for c in ("B17", "C17"): db[c].fill = fill(BG); db[c].border = Border()
cell(db, "B18", "Income", bold=True); cell(db, "C18", f"={INCOME}", CUR, bold=True, align="right")
cell(db, "B22", "Committed before you spend a cent", bold=True, color=ORANGE); cell(db, "C22", "=SUM(C19:C21)", CUR, bold=True, color=ORANGE, align="right")
cell(db, "B23", "Left for everything else", bold=True, color=GREEN); cell(db, "C23", "=C18-C22", CUR, bold=True, color=GREEN, align="right")
for r in (19, 20, 21):
    db[f"D{r}"] = f"=IF($C$18>0,C{r}/$C$18,0)"; db[f"D{r}"].number_format = '0%'; db[f"D{r}"].font = F(MUTED, False, 9)
db.column_dimensions["D"].width = 7
# chart helper (pie needs non-negative values)
db["K17"], db["L17"] = "Slice", "Amount"
for i, (lab, src) in enumerate([("Loans & cards", "C19"), ("Installments & BNPL", "C20"), ("Subscriptions", "C21"), ("Left over", "C23")]):
    db[f"K{18 + i}"] = lab; db[f"L{18 + i}"] = f"=MAX(0,{src})"
for r in range(17, 22):
    for c in "KL": db[f"{c}{r}"].font = F(BG, False, 8)
pie = PieChart(); pie.title = "Where each month's income goes"
pie.add_data(Reference(db, min_col=12, min_row=17, max_row=21), titles_from_data=True)
pie.set_categories(Reference(db, min_col=11, min_row=18, max_row=21))
pie.dataLabels = DataLabelList(); pie.dataLabels.showPercent = True
pie.height, pie.width = 8, 12; db.add_chart(pie, "E15")
bar = BarChart(); bar.type = "bar"; bar.title = "Real APR by debt"; bar.y_axis.numFmt = "0%"; bar.legend = None
bar.add_data(Reference(de, min_col=10, min_row=5, max_row=N0 + len(sample) - 1), titles_from_data=True)
bar.set_categories(Reference(de, min_col=2, min_row=N0, max_row=N0 + len(sample) - 1))
bar.height, bar.width = 8, 12; db.add_chart(bar, "E32")
# upcoming list
db["B26"] = "COMING UP (NEXT CHARGES)"; db["B26"].font = F(TEAL, True, 10)
header(db, 27, 2, ["Subscription", "Days to charge"], TEAL)
for k in range(1, 6):
    r = 27 + k
    cell(db, f"C{r}", f'=IFERROR(SMALL({sr("O")},{k}),"")', "0", align="center")
    cell(db, f"B{r}", f'=IF(C{r}="","",INDEX({sr("B")},MATCH(C{r},{sr("O")},0)))')
db["B34"] = "Two subscriptions due the same day show the first name twice. Check the Subscriptions tab."; db["B34"].font = F(MUTED, False, 8, True)
db["B36"] = "Estimates only, not financial advice. Real APR is worked out from your payments and excludes fees charged separately; your contract is the final word."
db["B36"].font = F(MUTED, False, 8, True)

# ---------------- How to use ----------------
hw = sheet("How to Use", {"A": 3, "B": 100})
title(hw, "How to use", "Five minutes to set up.")
steps = ["1. Settings: enter your monthly take-home income. Adjust the red/amber APR levels if you like (defaults 20% / 10%).",
         "2. Debts & Installments: one row per debt. BLUE cells are yours; ORANGE-header columns calculate themselves.",
         "   • Installment / BNPL / store financing: enter Original amount, Monthly payment, Total payments and Payments made. Leave Stated APR empty: the Real APR is worked out for you.",
         "   • Credit card / loan: enter Monthly payment, Stated APR and Balance now.",
         "   • Family loan with no interest: enter 0% as the Stated APR.",
         "3. Subscriptions: one row per subscription. Set Keep? to No on anything you'd cancel: the Dashboard shows the yearly saving.",
         "   • Free trial? Set Free trial? to Yes and enter Trial days. You get a 'cancel by' date and a countdown.",
         "4. Dashboard: see what share of your income is committed before you spend anything, your highest-cost debt, and what's coming up.",
         "",
         "Avalanche order = pay extra on #1 first (highest APR, saves the most interest). Snowball order = smallest balance first (quick wins).",
         "The example rows are made up. Replace or delete them.",
         "",
         "Estimates only. Not financial advice. The Real APR is an estimate from your payment schedule and does not include fees charged separately."]
for i, s in enumerate(steps):
    hw[f"B{5 + i}"] = s; hw[f"B{5 + i}"].font = F(TXT if s[:1].isdigit() else MUTED, s[:1].isdigit(), 10)
    hw[f"B{5 + i}"].alignment = Alignment(wrap_text=True, vertical="top")
wb.move_sheet("How to Use", -(len(wb.sheetnames) - 2))
for ws in wb.worksheets: ws.sheet_properties.tabColor = {"Dashboard": ORANGE, "How to Use": TEAL}.get(ws.title, "3A3F58")
wb.save("hidden-cost-tracker-prototype.xlsx"); print("saved", wb.sheetnames)
