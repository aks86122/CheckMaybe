"""Holiday Hidden Cost Tracker (dark green / gold / berry): gift list and budget, holiday payment plans
(pay-in-4, BNPL, store financing) with real APR and a month-by-month bill, and Black Friday free trials.

Builds Holiday-Hidden-Cost-Tracker.xlsx. For the Google Sheets version: import, File > Save as Google
Sheets, then run holiday_polish.gs once (checkboxes, filter bars, native dark charts, locale)."""
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties
from openpyxl.comments import Comment

BG, PANEL, PANEL2, INPUT, INPUT2, LINE, HEAD = "101A14", "17251D", "1B2B22", "1F3127", "25392D", "2F4838", "0B1310"
TXT, MUTED, GOLD, BERRY, GREEN, TEAL, ICE, PURPLE = "F3EFE4", "A7B2A3", "E0B05A", "E2687B", "74D39A", "5CC8B8", "9CC9F0", "B79CFF"
CHIP = {"gold": (GOLD, "3D3218"), "berry": (BERRY, "43202A"), "green": (GREEN, "1C3A2A"), "teal": (TEAL, "173833"),
        "ice": (ICE, "1D3045"), "purple": (PURPLE, "2E2546"), "grey": ("C2C8BE", "2A3530")}
F = lambda c=TXT, b=False, s=10, i=False, strike=False: Font(name="Arial", color=c, bold=b, size=s, italic=i, strike=strike)
fill = lambda c: PatternFill("solid", fgColor=c)
thin = Side(style="thin", color=LINE); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
CUR, PCT, DATE = '$#,##0.00;($#,##0.00);"-"', '0.0%;-0.0%;"-"', 'mmm d, yyyy'
APR = '0.0%'                         # 0% is a real answer for pay-in-4, so show it instead of "-"
DAYS = '[=0]"Today";[=1]"Tomorrow";0" days"'
IN_H, OUT_H = GOLD, BERRY            # header colours: you fill in / calculated
NOTE = "GOLD headers: you fill in   ·   BERRY headers: calculated for you"

wb = Workbook()
def sheet(name, widths, first=False, rows=40, cols=12):
    ws = wb.active if first else wb.create_sheet(name)
    ws.title = name; ws.sheet_view.showGridLines = False
    for col, w in widths.items(): ws.column_dimensions[col].width = w
    for r in range(1, rows + 1):
        for c in range(1, cols + 1): ws.cell(r, c).fill = fill(BG)
    return ws

def title(ws, text, sub=None):
    ws["B2"] = text; ws["B2"].font = F(TXT, True, 18); ws.row_dimensions[2].height = 30
    if sub: ws["B3"] = sub; ws["B3"].font = F(MUTED, False, 10, True)

def header(ws, row, col, labels, color=OUT_H):
    for i, l in enumerate(labels):
        c = ws.cell(row, col + i, l if str(l).startswith("=") else str(l).upper()); c.font = F(color, True, 8); c.fill = fill(HEAD)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(left=thin, right=thin, top=thin, bottom=Side(style="medium", color=color))
    ws.row_dimensions[row].height = 34

def cell(ws, ref, v, fmt=None, inp=False, bold=False, color=None, align="left", band=False, size=10):
    c = ws[ref]; c.value = v; c.border = BOX
    c.fill = fill((INPUT2 if band else INPUT) if inp else (PANEL2 if band else PANEL))
    c.font = F(color or TXT, bold, size); c.alignment = Alignment(horizontal=align, vertical="center")
    if fmt: c.number_format = fmt
    return c

def chips(ws, rng, mapping):
    for v, name in mapping.items():
        fg, bg = CHIP[name]
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{v}"'], fill=fill(bg), font=F(fg, True, 9)))

def listdv(ws, src, rng):
    v = DataValidation(type="list", formula1=src, allow_blank=True); ws.add_data_validation(v); v.add(rng)

d = lambda y, m, dd: dt.date(y, m, dd)

# ---------------- Settings ----------------
st = sheet("Settings", {"A": 3, "B": 30, "C": 18, "D": 15, "E": 17, "F": 13, "G": 11, "H": 44}, rows=26, cols=9)
title(st, "Settings", "Gold-header cells are yours to edit.")
header(st, 4, 2, ["Setting", "Value"], IN_H); header(st, 4, 4, ["What it does"], MUTED); st.merge_cells("D4:H4")
rows = [("Holiday budget", 1500, CUR, "Everything you plan to spend on gifts this season."),
        ("Monthly take-home income", 4200, CUR, "Income after tax, per month. Used for the share-of-income figures."),
        ("First month to plan", d(2026, 12, 1), "mmm yyyy", "The month-by-month bill starts here and runs 5 months."),
        ("High-cost APR (berry flag)", 0.20, PCT, "Any plan at or above this real APR is flagged HIGH."),
        ("Check APR (gold flag)", 0.10, PCT, "At or above this real APR is flagged CHECK."),
        ("Warn me this many days ahead", 7, "0", "Free trials ending within this many days are highlighted."),
        ("Today's date", "=TODAY()", DATE, "Updates automatically. Every date calculation uses it.")]
for i, (lab, val, fmt, note) in enumerate(rows):
    r = 5 + i
    cell(st, f"B{r}", lab); cell(st, f"C{r}", val, fmt, inp=(i < 6), align="right", color=GOLD if i < 6 else MUTED)
    st.merge_cells(f"D{r}:H{r}"); st[f"D{r}"] = note; st[f"D{r}"].font = F(MUTED, False, 9, True); st.row_dimensions[r].height = 22
BUDGET, INCOME, START, RED_APR, AMB_APR, WARN, TODAY = ("Settings!$C$5", "Settings!$C$6", "Settings!$C$7", "Settings!$C$8",
                                                        "Settings!$C$9", "Settings!$C$10", "Settings!$C$11")
header(st, 13, 2, ["Relationship", "Paid with", "Gift status", "Plan type", "Pays every", "Billing"], MUTED)
lists = {"B": ["Family", "Partner", "Kids", "Friend", "Coworker", "Other"],
         "C": ["Cash / debit", "Credit card", "Pay-in-4", "Monthly BNPL", "Store financing", "Gift card", "Other"],
         "D": ["Idea", "Bought", "Wrapped", "Given"],
         "E": ["Pay-in-4", "Monthly BNPL", "Store financing", "Credit card plan", "Other"],
         "F": ["2 weeks", "Month"],
         "G": ["Weekly", "Monthly", "Yearly"]}
for col, items in lists.items():
    for j in range(8):
        cell(st, f"{col}{14 + j}", items[j] if j < len(items) else None, inp=True, band=j % 2 == 1)
st["H14"] = "Rename or add items here and the drop-downs follow."; st["H14"].font = F(MUTED, False, 9, True)
st["H15"] = "Keep the Paid with names Pay-in-4, Monthly BNPL and Store financing:"; st["H15"].font = F(MUTED, False, 9, True)
st["H16"] = "the Gift List uses them to mark gifts that are on a payment plan."; st["H16"].font = F(MUTED, False, 9, True)
REL, PAID, STATUS, PLAN, EVERY, BILL = ("=Settings!$B$14:$B$21", "=Settings!$C$14:$C$21", "=Settings!$D$14:$D$21",
                                        "=Settings!$E$14:$E$21", "=Settings!$F$14:$F$15", "=Settings!$G$14:$G$16")

# ---------------- Gift List ----------------
gl = sheet("Gift List", {"A": 3, "B": 18, "C": 13, "D": 26, "E": 11, "F": 11, "G": 16, "H": 11, "I": 22, "J": 12, "K": 12, "L": 4},
           rows=57, cols=13)
title(gl, "Gift List")
gl["H2"] = NOTE; gl["H2"].font = F(MUTED, False, 8, True)
gl.row_dimensions[4].height = 40
header(gl, 5, 2, ["Recipient", "Relationship", "Gift", "Budget", "Actual cost", "Paid with", "Status", "Where / notes"], IN_H)
header(gl, 5, 10, ["Over / under", "Payment plan"], OUT_H)
gsample = [("Mom", "Family", "Cashmere scarf", 80, 74, "Credit card", "Bought", ""),
           ("Dad", "Family", "Smart watch", 250, 279, "Pay-in-4", "Bought", "Black Friday deal"),
           ("Alex", "Partner", "4K TV", 600, 649, "Store financing", "Bought", "12 monthly payments"),
           ("Sam", "Kids", "Game console", 450, 499, "Monthly BNPL", "Wrapped", ""),
           ("Jordan", "Friend", "Concert tickets", 120, 140, "Credit card", "Bought", ""),
           ("Grandma", "Family", "Photo book", 40, 38, "Cash / debit", "Wrapped", ""),
           ("Office gift swap", "Coworker", "Candle set", 25, 25, "Cash / debit", "Given", "$25 limit"),
           ("Teacher", "Other", "Coffee gift card", 30, 30, "Gift card", "Bought", ""),
           ("Sister", "Family", "Skincare set", 60, None, None, "Idea", ""),
           ("Nephew", "Kids", "Building block set", 60, None, None, "Idea", "")]
G0, G1 = 6, 55
for r in range(G0, G1 + 1):
    s = gsample[r - G0] if r - G0 < len(gsample) else (None,) * 8
    band = (r - G0) % 2 == 1
    for j, (col, fmt) in enumerate(zip("BCDEFGHI", [None, None, None, CUR, CUR, None, None, None])):
        cell(gl, f"{col}{r}", s[j], fmt, inp=True, band=band, align="right" if col in "EF" else "center" if col in "CGH" else "left")
    cell(gl, f"J{r}", f'=IF(AND(ISNUMBER($E{r}),ISNUMBER($F{r})),$E{r}-$F{r},"")', CUR, band=band, align="right")
    cell(gl, f"K{r}", f'=IF(OR($G{r}="Pay-in-4",$G{r}="Monthly BNPL",$G{r}="Store financing"),"On a plan","")', band=band, align="center")
    gl[f"L{r}"] = f'=IF(AND($H{r}="Idea",ISNUMBER($E{r})),$E{r}+ROW()/100000,"")'          # still-to-buy order
    gl.row_dimensions[r].height = 22
gl["L5"] = "key"; gl.column_dimensions["L"].hidden = True
gr = lambda c: f"'Gift List'!${c}${G0}:${c}${G1}"
SPENT = f'SUM({gr("F")})'
gl["B3"] = (f'=COUNTA({gr("B")})&" gifts   ·   spent "&TEXT({SPENT},"$#,##0")&" of "&TEXT({BUDGET},"$#,##0")'
            f'&"   ·   "&IF({BUDGET}-{SPENT}>=0,TEXT({BUDGET}-{SPENT},"$#,##0")&" left",TEXT({SPENT}-{BUDGET},"$#,##0")&" over budget")'
            f'&"   ·   "&COUNTIF({gr("K")},"On a plan")&" on a payment plan"')
gl["B3"].font = F(TXT, True, 10)
gl["K5"].comment = Comment("Gifts paid with Pay-in-4, Monthly BNPL or Store financing. Add each one to Holiday Payments to see what it costs month by month.", "CheckMaybe")
listdv(gl, REL, f"C{G0}:C{G1}"); listdv(gl, PAID, f"G{G0}:G{G1}"); listdv(gl, STATUS, f"H{G0}:H{G1}")
chips(gl, f"C{G0}:C{G1}", {"Family": "berry", "Partner": "purple", "Kids": "gold", "Friend": "teal", "Coworker": "ice", "Other": "grey"})
chips(gl, f"G{G0}:G{G1}", {"Cash / debit": "green", "Credit card": "ice", "Pay-in-4": "gold", "Monthly BNPL": "berry",
                            "Store financing": "berry", "Gift card": "teal", "Other": "grey"})
chips(gl, f"H{G0}:H{G1}", {"Idea": "grey", "Bought": "ice", "Wrapped": "gold", "Given": "green"})
chips(gl, f"K{G0}:K{G1}", {"On a plan": "berry"})
gl.conditional_formatting.add(f"J{G0}:J{G1}", FormulaRule(formula=[f'AND(ISNUMBER($J{G0}),$J{G0}<0)'], font=F(BERRY, True)))
gl.conditional_formatting.add(f"J{G0}:J{G1}", FormulaRule(formula=[f'AND(ISNUMBER($J{G0}),$J{G0}>0)'], font=F(GREEN, True)))
gl.auto_filter.ref = f"B5:K{G1}"

# ---------------- Holiday Payments ----------------
hp = sheet("Holiday Payments", {"A": 3, "B": 20, "C": 15, "D": 11, "E": 11, "F": 10, "G": 10, "H": 13, "I": 11, "J": 11,
                                "K": 12, "L": 11, "M": 10, "N": 13, "O": 10, "P": 12, "Q": 9, "R": 11, "S": 11, "T": 11,
                                "U": 11, "V": 11, "W": 4}, rows=27, cols=24)
title(hp, "Holiday Payments")
hp["N2"] = NOTE; hp["N2"].font = F(MUTED, False, 8, True)
hp.row_dimensions[4].height = 40
header(hp, 5, 2, ["Item", "Plan type", "Price", "Payment", "Number of payments", "Pays every", "First payment",
                  "1st paid at checkout?", "Stated APR (if known)"], IN_H)
header(hp, 5, 11, ["Total you'll pay", "Extra cost", "Real APR", "Last payment", "Payments left", "Still to pay", "Flag"], OUT_H)
header(hp, 5, 18, [f'=UPPER(TEXT(EDATE({START},{j}),"mmm yyyy"))' for j in range(5)], OUT_H)
psample = [("Smart watch", "Pay-in-4", 279, 69.75, 4, "2 weeks", d(2026, 11, 27), "Yes", None),
           ("Flights home", "Pay-in-4", 420, 105, 4, "2 weeks", d(2026, 12, 15), "Yes", None),
           ("4K TV", "Store financing", 649, 59, 12, "Month", d(2026, 12, 27), "No", None),
           ("Game console", "Monthly BNPL", 499, 89.5, 6, "Month", d(2026, 12, 27), "No", None)]
P0, P1 = 6, 25
for r in range(P0, P1 + 1):
    s = psample[r - P0] if r - P0 < len(psample) else (None,) * 9
    band = (r - P0) % 2 == 1
    for j, (col, fmt) in enumerate(zip("BCDEFGHIJ", [None, None, CUR, CUR, "0", None, DATE, None, PCT])):
        cell(hp, f"{col}{r}", s[j], fmt, inp=True, band=band, align="left" if col == "B" else "center" if col in "CGI" else "right")
    live = f'AND($B{r}<>"",N($E{r})>0,N($F{r})>0)'
    dated = f'AND({live},ISNUMBER($H{r}))'
    per = f'IF($G{r}="2 weeks",26,12)'
    mdiff = f'DATEDIF($H{r},{TODAY},"M")'
    past = (f'IF({TODAY}<=$H{r},0,MIN($F{r},IF($G{r}="2 weeks",INT(({TODAY}-$H{r}-1)/14)+1,'
            f'{mdiff}+IF(EDATE($H{r},{mdiff})<{TODAY},1,0))))')
    f = {"K": f'=IF({live},$E{r}*$F{r},"")',
         "L": f'=IF(AND(ISNUMBER($K{r}),ISNUMBER($D{r})),MAX(0,$K{r}-$D{r}),"")',
         "M": (f'=IF(NOT({live}),"",IF(ISNUMBER($J{r}),$J{r},IF(NOT(ISNUMBER($D{r})),"",IF($K{r}<=$D{r}+0.005,0,'
               f'IFERROR(RATE($F{r},-$E{r},$D{r},0,IF($I{r}="Yes",1,0))*{per},"")))))'),
         "N": f'=IF(NOT({dated}),"",IF($G{r}="2 weeks",$H{r}+14*($F{r}-1),EDATE($H{r},$F{r}-1)))',
         "O": f'=IF(NOT({dated}),"",$F{r}-{past})',
         "P": f'=IF(ISNUMBER($O{r}),$O{r}*$E{r},"")',
         "Q": f'=IF(NOT(ISNUMBER($M{r})),"",IF($M{r}>={RED_APR},"HIGH",IF($M{r}>={AMB_APR},"CHECK","OK")))'}
    for j, col in enumerate("RSTUV"):
        ms, me = f"EDATE({START},{j})", f"EOMONTH({START},{j})"
        bi = f'MAX(0,MIN($F{r}-1,INT(({me}-$H{r})/14))-MAX(0,-INT(-({ms}-$H{r})/14))+1)'
        dm = f'((YEAR({ms})*12+MONTH({ms}))-(YEAR($H{r})*12+MONTH($H{r})))'
        mo = f'IF(AND({dm}>=0,{dm}<$F{r}),1,0)'
        f[col] = f'=IF(NOT({dated}),"",$E{r}*IF($G{r}="2 weeks",{bi},{mo}))'
    for col, fx in f.items():
        fmt = {"K": CUR, "L": CUR, "M": APR, "N": DATE, "O": "0", "P": CUR}.get(col, CUR if col in "RSTUV" else None)
        cell(hp, f"{col}{r}", fx, fmt, band=band, align="center" if col in "OQ" else "right")
    hp[f"W{r}"] = f'=IF(ISNUMBER($M{r}),$M{r}+ROW()/10000000,"")'                                  # APR order
    hp.row_dimensions[r].height = 22
hp["W5"] = "key"; hp.column_dimensions["W"].hidden = True
pr = lambda c: f"'Holiday Payments'!${c}${P0}:${c}${P1}"
hp["B3"] = (f'=COUNT({pr("K")})&" plans   ·   "&TEXT(SUM({pr("P")}),"$#,##0")&" still to pay   ·   extra cost "&TEXT(SUM({pr("L")}),"$#,##0")'
            f'&"   ·   last payment "&IF(COUNT({pr("N")})=0,"-",TEXT(MAX({pr("N")}),"mmm d, yyyy"))')
hp["B3"].font = F(TXT, True, 10)
hp["M5"].comment = Comment("Stated APR if you entered one. Otherwise worked out with RATE() from price, payment and number of payments (first payment at checkout counts as paid up front). Fees charged separately are not included. Your agreement is the final word.", "CheckMaybe")
hp["R5"].comment = Comment("What each plan charges in that calendar month. The first month is set in Settings.", "CheckMaybe")
listdv(hp, PLAN, f"C{P0}:C{P1}"); listdv(hp, EVERY, f"G{P0}:G{P1}"); listdv(hp, '"Yes,No"', f"I{P0}:I{P1}")
chips(hp, f"C{P0}:C{P1}", {"Pay-in-4": "gold", "Monthly BNPL": "berry", "Store financing": "berry", "Credit card plan": "ice", "Other": "grey"})
chips(hp, f"G{P0}:G{P1}", {"2 weeks": "teal", "Month": "purple"})
chips(hp, f"I{P0}:I{P1}", {"Yes": "green", "No": "grey"})
chips(hp, f"Q{P0}:Q{P1}", {"HIGH": "berry", "CHECK": "gold", "OK": "green"})
hp.conditional_formatting.add(f"M{P0}:M{P1}", FormulaRule(formula=[f'AND(ISNUMBER($M{P0}),$M{P0}>={RED_APR})'], font=F(BERRY, True)))
hp.conditional_formatting.add(f"L{P0}:L{P1}", FormulaRule(formula=[f'AND(ISNUMBER($L{P0}),$L{P0}>0)'], font=F(BERRY, True)))
hp.auto_filter.ref = f"B5:Q{P1}"

# ---------------- Free Trials ----------------
ft = sheet("Free Trials", {"A": 3, "B": 22, "C": 13, "D": 10, "E": 11, "F": 11, "G": 11, "H": 13, "I": 11, "J": 13, "K": 13,
                           "L": 14, "M": 4}, rows=27, cols=14)
title(ft, "Free Trials")
ft["H2"] = NOTE; ft["H2"].font = F(MUTED, False, 8, True)
ft.row_dimensions[4].height = 40
header(ft, 5, 2, ["Service", "Started", "Trial days", "Then costs", "Billing", "Cancelled?"], IN_H)
header(ft, 5, 8, ["Cancel by", "Days left", "First charge", "Yearly cost if kept", "Alert"], OUT_H)
tsample = [("Streaming service", d(2026, 11, 27), 30, 15.99, "Monthly", "No"),
           ("Music app", d(2026, 11, 28), 90, 10.99, "Monthly", "No"),
           ("Shipping membership", d(2026, 11, 20), 30, 14.99, "Monthly", "No"),
           ("Photo editor", d(2026, 11, 27), 14, 9.99, "Monthly", "Yes"),
           ("Fitness app", d(2026, 12, 26), 7, 79.99, "Yearly", "No")]
T0, T1 = 6, 25
for r in range(T0, T1 + 1):
    s = tsample[r - T0] if r - T0 < len(tsample) else (None,) * 6
    band = (r - T0) % 2 == 1
    for j, (col, fmt) in enumerate(zip("BCDEFG", [None, DATE, "0", CUR, None, None])):
        cell(ft, f"{col}{r}", s[j], fmt, inp=True, band=band, align="left" if col == "B" else "center" if col in "FG" else "right")
    live = f'AND($B{r}<>"",ISNUMBER($C{r}),N($D{r})>0)'
    f = {"H": f'=IF({live},$C{r}+$D{r}-1,"")',
         "I": f'=IF(ISNUMBER($H{r}),$H{r}-{TODAY},"")',
         "J": f'=IF({live},$C{r}+$D{r},"")',
         "K": f'=IF(AND({live},ISNUMBER($E{r})),$E{r}*IF($F{r}="Weekly",52,IF($F{r}="Yearly",1,12)),"")',
         "L": (f'=IF(NOT({live}),"",IF($G{r}="Yes","Cancelled",IF($I{r}<0,"Trial ended",'
               f'IF($I{r}<={WARN},"Cancel soon?","OK"))))')}
    for col, fx in f.items():
        cell(ft, f"{col}{r}", fx, {"H": DATE, "I": DAYS, "J": DATE, "K": CUR}.get(col), band=band, align="center" if col in "IL" else "right")
    ft[f"M{r}"] = f'=IF(AND(ISNUMBER($I{r}),N($I{r})>=0,$G{r}<>"Yes"),$I{r}+ROW()/1000,"")'          # open trials order
    ft.row_dimensions[r].height = 22
ft["M5"] = "key"; ft.column_dimensions["M"].hidden = True
tr = lambda c: f"'Free Trials'!${c}${T0}:${c}${T1}"
ft["B3"] = (f'=COUNT({tr("M")})&" open trials   ·   "&TEXT(SUMIFS({tr("K")},{tr("G")},"<>Yes"),"$#,##0")&" a year if you keep them all'
            f'   ·   cancelled so far saves "&TEXT(SUMIFS({tr("K")},{tr("G")},"Yes"),"$#,##0")&" a year"')
ft["B3"].font = F(TXT, True, 10)
ft["G5"].comment = Comment("Tick it once you've cancelled. The trial drops off the Dashboard and counts as savings.", "CheckMaybe")
listdv(ft, BILL, f"F{T0}:F{T1}"); listdv(ft, '"Yes,No"', f"G{T0}:G{T1}")
chips(ft, f"F{T0}:F{T1}", {"Weekly": "berry", "Monthly": "purple", "Yearly": "gold"})
chips(ft, f"G{T0}:G{T1}", {"Yes": "green", "No": "grey"})
chips(ft, f"L{T0}:L{T1}", {"Cancel soon?": "berry", "Trial ended": "berry", "Cancelled": "green", "OK": "grey"})
ft.conditional_formatting.add(f"I{T0}:I{T1}", FormulaRule(formula=[f'AND(ISNUMBER($I{T0}),$I{T0}>=0,$I{T0}<={WARN})'], fill=fill(CHIP["berry"][1]), font=F(BERRY, True)))
ft.conditional_formatting.add(f"B{T0}:K{T1}", FormulaRule(formula=[f'$G{T0}="Yes"'], fill=fill(INPUT), font=F("6E7A70", False, 10, True, True)))
ft.auto_filter.ref = f"B5:L{T1}"

# ---------------- Chart Data ----------------
cd = sheet("Chart Data", {"A": 3, "B": 24, "C": 14, "D": 14, "E": 14, "F": 14}, rows=40, cols=7)
title(cd, "Chart Data", "Feeds the Dashboard charts and lists. Nothing to edit here.")
header(cd, 4, 2, ["Month", "Holiday payments due", "Share of income"], MUTED)
for j in range(5):
    r = 5 + j; col = "RSTUV"[j]
    cell(cd, f"B{r}", f'=TEXT(EDATE({START},{j}),"mmm yyyy")')
    cell(cd, f"C{r}", f'=SUM({pr(col)})', CUR, align="right")
    cell(cd, f"D{r}", f'=IF({INCOME}>0,C{r}/{INCOME},0)', PCT, align="right")
header(cd, 11, 2, ["How gifts were paid", "Spent"], MUTED)
for j in range(7):
    r = 12 + j; lab = f"Settings!$C${14 + j}"
    cell(cd, f"B{r}", f'=IF({lab}="","",{lab})')
    cell(cd, f"C{r}", f'=IF({lab}="",0,SUMIFS({gr("F")},{gr("G")},{lab}))', CUR, align="right")
header(cd, 21, 2, ["Plan (highest APR first)", "Real APR", "Flag", "Still to pay"], MUTED)
for i in range(1, 7):
    r = 21 + i; m = f'MATCH(LARGE({pr("W")},{i}),{pr("W")},0)'
    for col, src, fmt in (("B", "B", None), ("C", "M", APR), ("D", "Q", None), ("E", "P", CUR)):
        cell(cd, f"{col}{r}", f'=IFERROR(INDEX({pr(src)},{m}),"")', fmt, align="left" if col == "B" else "right")
header(cd, 29, 2, ["Still to buy (biggest budget first)", "Gift", "Budget", "Relationship"], MUTED)
for i in range(1, 6):
    r = 29 + i; m = f'MATCH(LARGE({gr("L")},{i}),{gr("L")},0)'
    for col, src, fmt in (("B", "B", None), ("C", "D", None), ("D", "E", CUR), ("E", "C", None)):
        cell(cd, f"{col}{r}", f'=IFERROR(INDEX({gr(src)},{m}),"")', fmt, align="right" if col == "D" else "left")

# ---------------- Dashboard ----------------
db = sheet("Dashboard", {"A": 2, "B": 24, "C": 14, "D": 14, "E": 14, "F": 16, "G": 16, "H": 16, "I": 18, "J": 2},
           first=True, rows=60, cols=11)
wb.move_sheet("Dashboard", -(len(wb.sheetnames) - 1))
db["B2"] = "HOLIDAY HIDDEN COST TRACKER"; db["B2"].font = F(GOLD, True, 22); db.row_dimensions[2].height = 34
db["B3"] = "Your gift budget, and what the holidays will still cost you in January."; db["B3"].font = F(MUTED, False, 10, True)
gap = Side(style="thick", color=BG)

def panel(r1, r2, c1, c2):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            x = db.cell(r, c); x.fill = fill(PANEL)
            x.border = Border(left=gap if c == c1 else None, right=gap if c == c2 else None,
                              top=gap if r == r1 else None, bottom=gap if r == r2 else None)

panel(2, 3, 8, 9)
db.merge_cells("H2:I2"); db.merge_cells("H3:I3")
db["H2"] = "TODAY"; db["H2"].font = F(MUTED, True, 8); db["H2"].alignment = Alignment(horizontal="center", vertical="bottom")
db["H3"] = f"={TODAY}"; db["H3"].number_format = "dddd, mmm d, yyyy"; db["H3"].font = F(TXT, True, 11)
db["H3"].alignment = Alignment(horizontal="center", vertical="top")

def kpi(r, c, label, fx, fmt, color, caption):
    panel(r, r + 2, c, c + 1)
    for k in range(3): db.merge_cells(start_row=r + k, start_column=c, end_row=r + k, end_column=c + 1)
    a, b, cap = db.cell(r, c), db.cell(r + 1, c), db.cell(r + 2, c)
    a.value = label; a.font = F(MUTED, True, 8); a.alignment = Alignment(indent=1, vertical="bottom")
    b.value = fx; b.font = F(color, True, 20); b.number_format = fmt; b.alignment = Alignment(horizontal="left", indent=1, vertical="center")
    cap.value = caption; cap.font = F(MUTED, False, 8, True); cap.alignment = Alignment(indent=1, vertical="top")
    db.row_dimensions[r].height = 20; db.row_dimensions[r + 1].height = 32; db.row_dimensions[r + 2].height = 20
    return b

done = f'COUNTIF({gr("H")},"Bought")+COUNTIF({gr("H")},"Wrapped")+COUNTIF({gr("H")},"Given")'
kpi(5, 2, "HOLIDAY BUDGET", f"={BUDGET}", CUR, TXT, "change it in Settings")
kpi(5, 4, "SPENT ON GIFTS", f"={SPENT}", CUR, GOLD, f'=IF({BUDGET}>0,TEXT({SPENT}/{BUDGET},"0%")&" of your budget","")')
left = kpi(5, 6, "LEFT IN BUDGET", f"={BUDGET}-{SPENT}", CUR, GREEN, f'=IF({BUDGET}-{SPENT}<0,"over budget","still to spend")')
db.conditional_formatting.add("F6", FormulaRule(formula=["F6<0"], font=F(BERRY, True, 20)))
kpi(5, 8, "GIFTS SORTED", f'={done}&" of "&COUNTA({gr("B")})', "General", TXT,
    f'=COUNTIF({gr("H")},"Idea")&" still ideas  ·  "&COUNTIF({gr("H")},"Given")&" given"')
mlabel = lambda j: f'="DUE IN "&UPPER(TEXT(EDATE({START},{j}),"mmmm"))'
mcap = lambda r: f'=TEXT(\'Chart Data\'!$D${r},"0%")&" of monthly take-home  ·  from holiday plans"'
kpi(9, 2, "", "='Chart Data'!$C$6", CUR, BERRY, mcap(6)); db["B9"] = mlabel(1)
kpi(9, 4, "", "='Chart Data'!$C$7", CUR, BERRY, mcap(7)); db["D9"] = mlabel(2)
kpi(9, 6, "HIGHEST REAL APR", f'=IF(COUNT({pr("M")})=0,"",MAX({pr("M")}))', APR, BERRY,
    f'=IF(COUNT({pr("M")})=0,"add a plan to see it",\'Chart Data\'!$B$22&"  ·  extra cost "&TEXT(SUM({pr("L")}),"$#,##0"))')
kpi(9, 8, "FREE TRIALS ENDING SOON", f'=COUNTIF({tr("L")},"Cancel soon?")', "0", GOLD,
    f'=IF(COUNT({tr("M")})=0,"no open free trials","first cancel-by: "&TEXT({TODAY}+INT(SMALL({tr("M")},1)),"mmm d"))')

def section(ref, text, color=TEAL):
    db[ref] = text; db[ref].font = F(color, True, 9); db[ref].alignment = Alignment(vertical="bottom")

def table(r0, labels, n, rowfx, fmts, aligns, color=TEAL):
    header(db, r0, 2, labels, color); db.row_dimensions[r0].height = 24
    for k in range(1, n + 1):
        r = r0 + k
        for j, (fx, fmt, al) in enumerate(zip(rowfx(k), fmts, aligns)):
            cell(db, db.cell(r, 2 + j).coordinate, fx, fmt, align=al, band=k % 2 == 0)
        db.row_dimensions[r].height = 21

section("B13", "HOLIDAY PAYMENTS  ·  MONTH BY MONTH", BERRY)
table(14, ["Month", "Due", "Share of income", "Running total"], 5,
      lambda k: [f"='Chart Data'!B{4 + k}", f"='Chart Data'!C{4 + k}", f"='Chart Data'!D{4 + k}", f"=SUM('Chart Data'!$C$5:C{4 + k})"],
      [None, CUR, PCT, CUR], ["left", "right", "right", "right"], BERRY)
db.conditional_formatting.add("D15:D19", FormulaRule(formula=["D15>=0.1"], fill=fill(CHIP["berry"][1]), font=F(BERRY, True)))
db.conditional_formatting.add("D15:D19", FormulaRule(formula=["D15>=0.05"], fill=fill(CHIP["gold"][1]), font=F(GOLD, True)))

pick = lambda key, k, col: f'IFERROR(INDEX({tr(col)},MATCH(SMALL({tr(key)},{k}),{tr(key)},0)),"")'
section("B21", "FREE TRIALS  ·  CANCEL BEFORE YOU'RE CHARGED", GOLD)
table(22, ["Free trial", "Then costs", "Cancel by", "Days left"], 4,
      lambda k: [f"={pick('M', k, 'B')}", f"={pick('M', k, 'E')}", f"={pick('M', k, 'H')}", f"={pick('M', k, 'I')}"],
      [None, CUR, "mmm d", DAYS], ["left", "right", "center", "center"], GOLD)
db.conditional_formatting.add("E23:E26", FormulaRule(formula=[f"AND(ISNUMBER(E23),E23<={WARN})"], fill=fill(CHIP["berry"][1]), font=F(BERRY, True)))

section("B28", "PAYMENT PLANS  ·  HIGHEST REAL APR FIRST", BERRY)
table(29, ["Plan", "Real APR", "Flag", "Still to pay"], 6,
      lambda k: [f"='Chart Data'!B{21 + k}", f"='Chart Data'!C{21 + k}", f"='Chart Data'!D{21 + k}", f"='Chart Data'!E{21 + k}"],
      [None, APR, None, CUR], ["left", "right", "center", "right"], BERRY)
chips(db, "D30:D35", {"HIGH": "berry", "CHECK": "gold", "OK": "green"})

section("B37", "STILL TO BUY", TEAL)
table(38, ["Recipient", "Gift", "Budget", "Relationship"], 5,
      lambda k: [f"='Chart Data'!B{29 + k}", f"='Chart Data'!C{29 + k}", f"='Chart Data'!D{29 + k}", f"='Chart Data'!E{29 + k}"],
      [None, None, CUR, None], ["left", "left", "right", "center"], TEAL)

section("B45", "WHAT THE FLAGS MEAN", MUTED)
flags = [("HIGH", "berry", f'="Real APR at or above "&TEXT({RED_APR},"0%")&". The monthly price hides a real cost."'),
         ("CHECK", "gold", f'="Real APR at or above "&TEXT({AMB_APR},"0%")&". Worth a second look."'),
         ("OK", "green", '="Little or no extra cost. Still a bill to plan for."')]
for i, (lab, col, fx) in enumerate(flags):
    r = 46 + i; fg, bg = CHIP[col]
    c = db[f"B{r}"]; c.value = lab; c.font = F(fg, True, 9); c.fill = fill(bg); c.alignment = Alignment(horizontal="center", vertical="center")
    db[f"C{r}"] = fx; db[f"C{r}"].font = F(MUTED, False, 9); db[f"C{r}"].alignment = Alignment(indent=1, vertical="center")
    db.row_dimensions[r].height = 20
db["B50"] = ("Pay-in-4 with no fee and the first payment at checkout works out to 0%. "
             "0% store offers can charge back-dated interest if a payment is late: check your agreement.")
db["B50"].font = F(MUTED, False, 8, True)

section("F13", "HOLIDAY PAYMENTS BY MONTH", BERRY)
section("F29", "HOW YOU PAID FOR GIFTS", GOLD)
db["B56"] = ("Estimates only, not financial advice. Real APR is worked out from your payments and excludes fees charged "
             "separately; your agreement is the final word. Example rows are made up.")
db["B56"].font = F(MUTED, False, 8, True)

def rich(color, sz=900, bold=False):
    cp = CharacterProperties(sz=sz, b=bold, solidFill=color)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])

def dark(ch, axes=True):
    ch.graphical_properties = GraphicalProperties(solidFill=PANEL, ln=LineProperties(noFill=True))
    ch.plot_area.graphicalProperties = GraphicalProperties(noFill=True, ln=LineProperties(noFill=True))
    if ch.legend is not None: ch.legend.txPr = rich(TXT)
    ch.visible_cells_only = False
    if axes:
        for ax in (ch.x_axis, ch.y_axis):
            ax.delete = False; ax.txPr = rich(MUTED, 800)
            ax.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=LINE))
        ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=LINE))

CDS = wb["Chart Data"]
col = BarChart(); col.type = "col"; col.gapWidth = 70; col.legend = None
col.add_data(Reference(CDS, min_col=3, min_row=4, max_row=9), titles_from_data=True)
col.set_categories(Reference(CDS, min_col=2, min_row=5, max_row=9))
col.series[0].graphicalProperties = GraphicalProperties(solidFill=BERRY, ln=LineProperties(noFill=True))
col.y_axis.number_format = '$#,##0'; dark(col)
col.width, col.height = 12.9, 7.0; db.add_chart(col, "F14")

don = DoughnutChart(holeSize=58)
don.add_data(Reference(CDS, min_col=3, min_row=11, max_row=18), titles_from_data=True)
don.set_categories(Reference(CDS, min_col=2, min_row=12, max_row=18))
for i, colr in enumerate([GREEN, ICE, GOLD, BERRY, "A8384B", TEAL, "8A948A"]):
    don.series[0].dPt.append(DataPoint(idx=i, spPr=GraphicalProperties(solidFill=colr, ln=LineProperties(solidFill=PANEL))))
don.dataLabels = DataLabelList(); don.dataLabels.showPercent = True; don.dataLabels.showVal = False
don.dataLabels.showCatName = False; don.dataLabels.showSerName = False; don.dataLabels.showLegendKey = False
don.dataLabels.txPr = rich(BG, 900, True)
don.legend.position = "r"; dark(don, axes=False)
don.width, don.height = 12.9, 7.0; db.add_chart(don, "F30")

# ---------------- How to use ----------------
hw = sheet("How to Use", {"A": 3, "B": 110}, rows=26, cols=3)
title(hw, "How to use", "Ten minutes to set up.")
steps = ["1. Settings: enter your holiday budget, monthly take-home income and the first month you want to plan (usually December).",
         "2. Gift List: one row per person. Gold-header columns are yours; berry-header columns calculate themselves.",
         "   • Set Paid with. Gifts on Pay-in-4, Monthly BNPL or Store financing are marked 'On a plan'.",
         "3. Holiday Payments: one row per payment plan (gifts, travel, anything bought on a plan this season).",
         "   • Enter the price, the payment, how many payments, how often and the first payment date.",
         "   • First payment taken at checkout (most pay-in-4)? Set '1st paid at checkout?' to Yes.",
         "   • You get the real APR, the extra cost, and what each plan charges in each month.",
         "4. Free Trials: one row per trial you start (Black Friday deals often come with one). Tick Cancelled? once you've cancelled.",
         "5. Dashboard: budget left, what your holiday plans add to January and February, the costliest plan, and trials to cancel.",
         "   • Use the filter buttons on each table's header row (or the filter bars above it in Google Sheets).",
         "",
         "The example rows are made up. Replace or delete them.",
         "",
         "Estimates only. Not financial advice. Real APR is an estimate from your payment schedule and does not include fees charged separately."]
for i, s in enumerate(steps):
    hw[f"B{5 + i}"] = s; hw[f"B{5 + i}"].font = F(TXT if s[:1].isdigit() else MUTED, s[:1].isdigit(), 10)
    hw[f"B{5 + i}"].alignment = Alignment(wrap_text=True, vertical="top")

order = ["Dashboard", "How to Use", "Settings", "Gift List", "Holiday Payments", "Free Trials", "Chart Data"]
wb._sheets = [wb[n] for n in order]
tabs = {"Dashboard": GOLD, "How to Use": TEAL, "Chart Data": "26302A", "Gift List": BERRY, "Holiday Payments": BERRY}
for ws in wb.worksheets: ws.sheet_properties.tabColor = tabs.get(ws.title, "2F4838")
wb.active = 0
wb.save("Holiday-Hidden-Cost-Tracker.xlsx"); print("saved", wb.sheetnames)
