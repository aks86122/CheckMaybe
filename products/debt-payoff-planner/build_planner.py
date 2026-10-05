"""Debt Payoff Planner (dark theme, matches Hidden Cost Tracker).

Beyond snowball/avalanche: early-payoff fees, lock-in periods, "not sure, ask your lender",
0%-promo / deferred-interest deadlines, behind-on-payments first, and a cash-flow order that frees
monthly money fastest (logic adapted from the founder's own debt planner). Every optional field can
be left blank without breaking the maths. Any currency.

Builds Debt-Payoff-Planner.xlsx. Each strategy runs a 360-month simulation on its own hidden sheet.
"""
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

BG, PANEL, PANEL2, INPUT, INPUT2, LINE, HEAD = "15171F", "1E2130", "232739", "262A3D", "2C3147", "33384F", "10121A"
TXT, MUTED, BLUE, ORANGE, RED, AMBER, GREEN, TEAL, PURPLE = "EDEBF5", "9A98AE", "8AB4FF", "F08A4B", "FF6B6B", "F2B544", "5FD38D", "4FC3C7", "B18CFF"
CHIP = {"purple": (PURPLE, "2E2250"), "teal": (TEAL, "16373A"), "amber": (AMBER, "4A3A15"), "green": (GREEN, "173A28"),
        "red": (RED, "4A1D22"), "blue": (BLUE, "1C2B4A"), "orange": (ORANGE, "4A2A18"), "grey": ("B8B6C8", "2F3242")}
F = lambda c=TXT, b=False, s=10, i=False: Font(name="Arial", color=c, bold=b, size=s, italic=i)
# Conditional-format fonts: colour/bold only. Font name or size in a dxf makes desktop Excel 'repair' styles.xml and drop all formatting.
FD = lambda c=TXT, b=False, *_a, **_k: Font(color=c, bold=b)
fill = lambda c: PatternFill("solid", fgColor=c)
thin = Side(style="thin", color=LINE); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
NUM, PCT, MON = '#,##0.00;(#,##0.00);"-"', '0.0%;-0.0%;"-"', 'mmm yyyy'

N = 12                    # debt rows
R0 = 7                    # first debt row on My Debts
RN = R0 + N - 1
MONTHS = 360
FAR = "DATE(2199,12,1)"
RULES = ["Allowed, no fee", "Allowed with a fee", "Locked until a date", "Not sure - ask lender"]
TYPES = ["Credit card", "Store card", "Personal loan", "Car loan", "Student loan", "Mortgage", "Buy now, pay later",
         "Medical bill", "Family / friend", "Other"]
STRATS = ["Smart (CheckMaybe order)", "Avalanche (highest APR first)", "Snowball (smallest balance first)",
          "Cash-flow (free up monthly money first)"]
CALC = ["Calc Smart", "Calc Avalanche", "Calc Snowball", "Calc Cashflow"]

wb = Workbook()
def sheet(name, widths, first=False, rows=40, cols=12):
    ws = wb.active if first else wb.create_sheet(name)
    ws.title = name; ws.sheet_view.showGridLines = False
    for col, w in widths.items(): ws.column_dimensions[col].width = w
    for r in range(1, rows + 1):
        for c in range(1, cols + 1): ws.cell(r, c).fill = fill(BG)
    return ws

def title(ws, text, sub=None):
    ws["B2"] = text; ws["B2"].font = F(TXT, True, 18); ws.row_dimensions[2].height = 30
    if sub: ws["B3"] = sub; ws["B3"].font = F(MUTED, False, 10, True)

def header(ws, row, col, labels, color=ORANGE, h=34):
    for i, l in enumerate(labels):
        c = ws.cell(row, col + i, l.upper()); c.font = F(color, True, 8); c.fill = fill(HEAD)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(left=thin, right=thin, top=thin, bottom=Side(style="medium", color=color))
    ws.row_dimensions[row].height = h

def cell(ws, ref, v, fmt=None, inp=False, bold=False, color=None, align="left", band=False, size=10, wrap=False):
    c = ws[ref]; c.value = v; c.border = BOX
    c.fill = fill((INPUT2 if band else INPUT) if inp else (PANEL2 if band else PANEL))
    c.font = F(color or TXT, bold, size); c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if fmt: c.number_format = fmt
    return c

def chips(ws, rng, mapping):
    for v, name in mapping.items():
        fg, bg = CHIP[name]
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{v}"'], fill=fill(bg), font=FD(fg, True, 9)))

def listdv(ws, items, rng):
    v = DataValidation(type="list", formula1='"' + ",".join(items) + '"', allow_blank=True); ws.add_data_validation(v); v.add(rng)

# ======================= Plan (settings + results) =======================
pl = sheet("Plan", {"A": 3, "B": 30, "C": 22, "D": 22, "E": 22, "F": 22, "G": 3, "H": 26, "I": 34}, first=True, rows=60, cols=10)
title(pl, "Debt Payoff Planner", "Set your budget here, list your debts on My Debts. Blue headers = you fill in.")
header(pl, 4, 2, ["Setting", "Value"], BLUE, 24)
S = {}
settings = [("budget", "Money for debts each month", 1500, NUM, "Total you can put toward ALL debts each month, minimums included."),
            ("start", "Plan starts", dt.date(2026, 11, 1), MON, "First month of the plan. Use the 1st of the month."),
            ("lump", "One-time extra (optional)", 0, NUM, "A bonus, tax refund or savings you'll throw at debt once. Leave 0 if none."),
            ("lumpm", "…in plan month #", 1, "0", "Which month the one-time extra goes in (1 = first month)."),
            ("strat", "Payoff order", STRATS[0], None, "Smart is the default. Compare all four on the right."),
            ("income", "Monthly take-home income (optional)", 3800, NUM, "Used only to spot money pressure: if minimums take 35%+ of income, Smart frees up monthly money first.")]
for i, (k, lab, v, fmt, note) in enumerate(settings):
    r = 5 + i; S[k] = f"Plan!$C${r}"
    cell(pl, f"B{r}", lab, band=i % 2); c = cell(pl, f"C{r}", v, fmt, inp=True, band=i % 2, color=BLUE, align="right")
    pl.merge_cells(f"D{r}:F{r}"); cell(pl, f"D{r}", note, color=MUTED, size=8, wrap=True, band=i % 2); pl.row_dimensions[r].height = 30
listdv(pl, STRATS, "C9")
pl["C9"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
# helpers (hidden column K)
pl["K5"] = f'=MATCH($C$9,{{"{STRATS[0]}","{STRATS[1]}","{STRATS[2]}","{STRATS[3]}"}},0)'; S["idx"] = "Plan!$K$5"
pl["K6"] = f"=SUMPRODUCT(('My Debts'!$D${R0}:$D${RN}>0)*N(+'My Debts'!$F${R0}:$F${RN}))"; S["minsum"] = "Plan!$K$6"
pl["K7"] = f"=AND(N($C$10)>0,$K$6/MAX(1,N($C$10))>=0.35)"; S["pressure"] = "Plan!$K$7"
pl["K8"] = f"=COUNTIFS('My Debts'!$D${R0}:$D${RN},\">0\")"; S["count"] = "Plan!$K$8"
pl.column_dimensions["K"].hidden = True

# Results: comparison table
pl["B13"] = "COMPARE THE FOUR ORDERS"; pl["B13"].font = F(TEAL, True, 9)
header(pl, 14, 2, ["", "Smart", "Avalanche", "Snowball", "Cash-flow"], TEAL, 24)
rows = [("Debt-free", MON, lambda s: f"='{s}'!$B$5"),
        ("Months to debt-free", "0", lambda s: f"='{s}'!$B$6"),
        ("Total interest", NUM, lambda s: f"='{s}'!$B$7"),
        ("Early-payoff fees paid", NUM, lambda s: f"='{s}'!$B$8"),
        ("Interest + fees", NUM, lambda s: f"='{s}'!$B$7+'{s}'!$B$8"),
        ("First debt cleared", MON, lambda s: f"='{s}'!$B$9")]
for i, (lab, fmt, fx) in enumerate(rows):
    r = 15 + i; cell(pl, f"B{r}", lab, band=i % 2, bold=True)
    for j, s in enumerate(CALC): cell(pl, f"{L(3 + j)}{r}", fx(s), fmt, band=i % 2, align="right")
pl.conditional_formatting.add("C19:F19", FormulaRule(formula=["AND(ISNUMBER(C19),C19=MIN($C$19:$F$19))"], font=FD(GREEN, True)))
pl.conditional_formatting.add("C15:F15", FormulaRule(formula=["AND(ISNUMBER(C15),C15=MIN($C$15:$F$15))"], font=FD(GREEN, True)))
pl["B21"] = "Green = best in that row. If a month shows \"-\", that plan doesn't finish within 30 years: raise the monthly amount."
pl["B21"].font = F(MUTED, False, 8, True)

# KPI cards for chosen order (H/I)
def kpi(r, label, fx, fmt, color):
    pl[f"H{r}"] = label; pl[f"H{r}"].font = F(MUTED, True, 8); pl[f"H{r}"].fill = fill(PANEL)
    pl[f"I{r}"] = fx; pl[f"I{r}"].font = F(color, True, 16); pl[f"I{r}"].number_format = fmt; pl[f"I{r}"].fill = fill(PANEL)
    pl[f"I{r}"].alignment = Alignment(horizontal="right"); pl.row_dimensions[r].height = 26
pl["H4"] = "YOUR CHOSEN ORDER"; pl["H4"].font = F(ORANGE, True, 9)
ch = lambda row: f"=INDEX(C{row}:F{row},{S['idx']})"
kpi(5, "DEBT-FREE", ch(15), MON, GREEN)
kpi(6, "MONTHS TO GO", ch(16), "0", TXT)
kpi(7, "INTEREST YOU'LL PAY", ch(17), NUM, RED)
kpi(8, "EARLY-PAYOFF FEES", ch(18), NUM, AMBER)
kpi(9, "MINIMUMS ONLY: INTEREST", f"=IF(COUNT('My Debts'!$P${R0}:$P${RN})=0,\"\",SUM('My Debts'!$P${R0}:$P${RN}))", NUM, MUTED)
kpi(10, "YOU SAVE VS MINIMUMS ONLY", f"=IF(OR(I9=\"\",NOT(ISNUMBER(I7))),\"\",MAX(0,I9-I7-I8))", NUM, GREEN)
pl["H11"] = "Minimums-only ignores debts that minimums never clear."; pl["H11"].font = F(MUTED, False, 8, True)

# warnings
pl["B23"] = "CHECKS"; pl["B23"].font = F(AMBER, True, 9)
warns = [f'=IF({S["budget"]}<{S["minsum"]},"⚠ Your monthly amount is less than your minimum payments ("&TEXT({S["minsum"]},"#,##0")&"). The plan assumes you still pay every minimum.","✓ Monthly amount covers all minimums ("&TEXT({S["minsum"]},"#,##0")&"). Extra each month: "&TEXT({S["budget"]}-{S["minsum"]},"#,##0"))',
         f"=IF(COUNTIF('My Debts'!$Q${R0}:$Q${RN},\"Minimum doesn't cover interest*\")>0,\"⚠ At least one minimum payment doesn't cover its interest: that debt grows unless you pay more.\",\"✓ Every minimum payment covers its interest.\")",
         f"=IF(COUNTIF('My Debts'!$Q${R0}:$Q${RN},\"Ask your lender*\")+COUNTIF('My Debts'!$Q${R0}:$Q${RN},\"Locked*\")+COUNTIF('My Debts'!$Q${R0}:$Q${RN},\"Wait*\")>0,\"ℹ Some debts get minimums only until they unlock (see My Debts › Early payoff check).\",\"✓ No locked or uncertain debts.\")",
         f'=IF({S["pressure"]},"ℹ Minimums take 35%+ of your income: Smart frees up monthly money first, then switches to highest APR.","ℹ Smart: behind-on-payments first, then 0%-promo deadlines, then highest APR.")']
warns.append(f"=IF(COUNTIF('My Debts'!$AE${R0}:$AE${RN},TRUE)>0,\"⚠ A deferred-interest debt is cleared after its promo ends in this order: back-interest would be added (not in these totals). Smart avoids this when it can.\",\"✓ No deferred-interest deadline missed in this order.\")")
for i, w in enumerate(warns):
    pl.merge_cells(f"B{24 + i}:F{24 + i}"); cell(pl, f"B{24 + i}", w, color=TXT, size=9, wrap=True, band=i % 2); pl.row_dimensions[24 + i].height = 28

# Priority list
pl["B29"] = "YOUR PAYOFF ORDER (CHOSEN)"; pl["B29"].font = F(ORANGE, True, 9); pl.row_dimensions[29].height = 22
header(pl, 30, 2, ["#", "Debt", "Balance", "APR", "Paid off", "Why / note"], ORANGE, 24)
pl.merge_cells("G30:I30")
for k in range(1, N + 1):
    r = 30 + k; b = k % 2 == 0
    m = f"MATCH({k},'My Debts'!$S${R0}:$S${RN},0)"
    cell(pl, f"B{r}", f'=IF({k}<={S["count"]},{k},"")', "0", band=b, align="center")
    cell(pl, f"C{r}", f"=IFERROR(INDEX('My Debts'!$B${R0}:$B${RN},{m}),\"\")", band=b)
    cell(pl, f"D{r}", f"=IFERROR(INDEX('My Debts'!$D${R0}:$D${RN},{m}),\"\")", NUM, band=b, align="right")
    cell(pl, f"E{r}", f"=IFERROR(INDEX('My Debts'!$E${R0}:$E${RN},{m}),\"\")", PCT, band=b, align="right")
    cell(pl, f"F{r}", f"=IFERROR(INDEX('My Debts'!$T${R0}:$T${RN},{m}),\"\")", MON, band=b, align="right")
    pl.merge_cells(f"G{r}:I{r}")
    cell(pl, f"G{r}", f"=IFERROR(INDEX('My Debts'!$V${R0}:$V${RN},{m}),\"\")", band=b, color=MUTED, size=9)
pl["B44"] = ("Estimates only, not financial advice. Interest is worked out monthly at APR÷12; your lender's method, fees and rules "
             "are the final word. Example debts are made up: replace them with yours.")
pl["B44"].font = F(MUTED, False, 8, True)

# ======================= My Debts =======================
md = sheet("My Debts", {"A": 3, "B": 22, "C": 16, "D": 14, "E": 9, "F": 13, "G": 11, "H": 20, "I": 12, "J": 13, "K": 13, "L": 11,
                       "M": 11, "N": 20, "O": 13, "P": 15, "Q": 38, "R": 12, "S": 9, "T": 12, "V": 40}, rows=RN + 6, cols=30)
title(md, "My Debts", "One row per debt. Blue = you fill in (blank is fine for optional ones). Orange = calculated.")
md["B4"] = "Any currency works: just use the same one everywhere. APR as a percent (24.99%)."; md["B4"].font = F(MUTED, False, 9, True)
yours = ["Debt name", "Type", "Balance now", "APR", "Minimum monthly payment", "Behind on payments?", "Early payoff rule",
         "Fee to pay off early", "Fee / lock ends", "0% / promo rate ends", "APR after promo", "Deferred interest?", "Notes"]
calc = ["Interest this month", "Interest if minimums only", "Early payoff check", "Extra can start", "Order #", "Paid off"]
header(md, 6, 2, yours, BLUE, 44); header(md, 6, 15, calc, ORANGE, 44)
md["V6"] = "WHY / NOTE"; md["V6"].font = F(ORANGE, True, 8); md["V6"].fill = fill(HEAD)
for c in "UWXYZ": md.column_dimensions[c].hidden = True
for c in ["AA", "AB", "AC", "AD", "AE"]: md.column_dimensions[c].hidden = True
tips = {"F6": "The smallest payment your lender asks for each month.",
        "G6": "Yes = you've missed payments or it's in collections. Smart puts these first.",
        "H6": "Leave blank if you don't know the rule yet: it's treated as 'Allowed, no fee'.",
        "I6": "The fee your contract charges for paying the whole balance early. Blank = no fee.",
        "J6": "When the early-payoff fee or the lock-in period ends. Blank on a lock = treated as locked for now.",
        "K6": "Only for 0% or promo-rate deals. Blank = no promo.",
        "L6": "The rate after the promo ends. Blank = same as APR.",
        "M6": "Yes = if not paid in full by the promo end, back-interest from day one gets added (common on US store cards)."}
for ref, t in tips.items(): md[ref].comment = Comment(t, "CheckMaybe")

sample = [("Store card (sample)", "Store card", 1200, 0, 60, "No", "Allowed, no fee", None, None, dt.date(2027, 5, 1), 0.2999, "Yes", "0% until May"),
          ("Credit card (sample)", "Credit card", 4800, 0.2449, 120, "No", "Allowed, no fee", None, None, None, None, None, ""),
          ("Car loan (sample)", "Car loan", 9500, 0.089, 310, "No", "Allowed with a fee", 400, dt.date(2027, 8, 1), None, None, None, "fee until Aug 2027"),
          ("Personal loan (sample)", "Personal loan", 3000, 0.159, 140, "No", "Not sure - ask lender", None, None, None, None, None, "call lender"),
          ("Medical bill (sample)", "Medical bill", 650, 0, 50, "Yes", "", None, None, None, None, None, "late notice")]
P = lambda col, r: f"{col}{r}"
for i in range(N):
    r = R0 + i; b = i % 2 == 1
    vals = sample[i] if i < len(sample) else [None] * 13
    fmts = [None, None, NUM, PCT, NUM, None, None, NUM, MON, MON, PCT, None, None]
    for j, (v, fm) in enumerate(zip(vals, fmts)):
        cell(md, f"{L(2 + j)}{r}", v if v != "" else None, fm, inp=True, band=b, align="right" if fm else "left")
    act = f"AND(ISNUMBER($D{r}),$D{r}>0)"
    apr = f"N(+$E{r})"
    minp = f"N(+$F{r})"
    # O interest this month
    cell(md, f"O{r}", f'=IF({act},$D{r}*{apr}/12,"")', NUM, band=b, align="right")
    # P interest if minimums only (current APR)
    cell(md, f"P{r}", f'=IF(NOT({act}),"",IF({minp}<=0,"No minimum",IF({minp}<=$D{r}*{apr}/12,"Never at minimum",IF({apr}=0,0,{minp}*NPER({apr}/12,-{minp},$D{r})-$D{r}))))', NUM, band=b, align="right")
    # Q early payoff check
    rule = f"$H{r}"
    q = (f'=IF(NOT({act}),"",IF(AND({minp}>0,{minp}<=$D{r}*{apr}/12),"Minimum doesn\'t cover interest: pay more here",'
         f'IF({minp}<=0,"Add the minimum payment",'
         f'IF({rule}="{RULES[3]}","Ask your lender before paying extra",'
         f'IF({rule}="{RULES[2]}",IF(N(+$J{r})=0,"Locked: add the unlock date","Locked until "&TEXT($J{r},"mmm yyyy")),'
         f'IF({rule}="{RULES[1]}",IF(N(+$I{r})=0,"Add the fee amount (blank = no fee)",'
         f'IF(NOT(ISNUMBER($P{r})),"Fee is worth paying: minimums never clear this",'
         f'IF(N(+$I{r})<$P{r},"Worth it: saves about "&TEXT($P{r}-$I{r},"#,##0")&" after the fee","Wait: the fee is more than the interest you\'d save"&IF(N(+$J{r})>0," (until "&TEXT($J{r},"mmm yyyy")&")","")))),'
         f'"OK to pay early"))))))')
    cell(md, f"Q{r}", q, band=b, color=TXT, size=9)
    # R extra can start (date)
    rr = (f'=IF(NOT({act}),"",IF({rule}="{RULES[3]}",{FAR},IF({rule}="{RULES[2]}",IF(N(+$J{r})=0,{FAR},$J{r}),'
          f'IF(AND({rule}="{RULES[1]}",N(+$I{r})>0,ISNUMBER($P{r}),N(+$I{r})>=N(+$P{r})),IF(N(+$J{r})=0,{FAR},$J{r}),DATE(1900,1,1)))))')
    md[f"U{r}"] = rr
    cell(md, f"R{r}", f'=IF($U{r}="","",IF($U{r}<={S["start"]},"Now",IF($U{r}>={FAR},"Not yet",TEXT($U{r},"mmm yyyy"))))', band=b, align="center")
    # hidden sort keys W..Z (smaller = earlier). Smart tiers: 0 behind, 1 deferred-interest promo still running, 2 rest.
    av = f"1000-{apr}*100+ROW()/1000000"
    cf = f"1000-IF({minp}>0,{minp}/$D{r},0)*100+ROW()/1000000"
    tier = f'IF($G{r}="Yes",0,IF(AND($M{r}="Yes",N(+$K{r})>{S["start"]}),1,2))'
    md[f"W{r}"] = f'=IF({act},{av},"")'
    md[f"X{r}"] = f'=IF({act},$D{r}+ROW()/1000000,"")'
    md[f"Y{r}"] = f'=IF({act},{cf},"")'
    md[f"Z{r}"] = f'=IF({act},{tier}*1000000+IF({tier}=0,{cf},IF({tier}=1,N(+$K{r})/100,IF({S["pressure"]},{cf},{av}))),"")'
    for k, c in zip(["AA", "AB", "AC", "AD"], "ZWXY"):  # ranks in Smart, Avalanche, Snowball, Cashflow order
        md[f"{k}{r}"] = f'=IF({c}{r}="","",COUNTIF({c}${R0}:{c}${RN},"<"&{c}{r})+1)'
    cell(md, f"S{r}", f'=IF({act},CHOOSE({S["idx"]},AA{r},AB{r},AC{r},AD{r}),"")', "0", band=b, align="center", bold=True, color=ORANGE)
    col = L(3 + i)
    cell(md, f"T{r}", f"=IF({act},CHOOSE({S['idx']}," + ",".join(f"'{s}'!{col}$10" for s in CALC) + "),\"\")", MON, band=b, align="right")
    md[f"AE{r}"] = f'=AND({act},$M{r}="Yes",N(+$K{r})>0,ISNUMBER($T{r}),N(+$T{r})>=N(+$K{r}))'
    why = (f'=IF(NOT({act}),"",IF($AE{r},"⚠ Cleared after the promo ends ("&TEXT($K{r},"mmm yyyy")&"): back-interest likely, not counted here. Try Smart.",IF($G{r}="Yes","Behind on payments: deal with this first",'
           f'IF(AND($M{r}="Yes",N(+$K{r})>{S["start"]}),"Deferred interest: clear by "&TEXT($K{r},"mmm yyyy")&" or back-interest is added",'
           f'IF($U{r}>{S["start"]},$Q{r},CHOOSE({S["idx"]},IF({S["pressure"]},"Frees up monthly money fastest","Highest APR among the rest"),'
           f'"APR "&TEXT({apr},"0.0%"),"Balance "&TEXT($D{r},"#,##0"),"Frees "&TEXT({minp},"#,##0")&"/month when cleared"))))))')
    cell(md, f"V{r}", why, band=b, color=MUTED, size=9)
listdv(md, RULES, f"H{R0}:H{RN}"); listdv(md, TYPES, f"C{R0}:C{RN}")
listdv(md, ["Yes", "No"], f"G{R0}:G{RN}"); listdv(md, ["Yes", "No"], f"M{R0}:M{RN}")
chips(md, f"G{R0}:G{RN}", {"Yes": "red", "No": "grey"}); chips(md, f"M{R0}:M{RN}", {"Yes": "amber", "No": "grey"})
chips(md, f"H{R0}:H{RN}", {RULES[0]: "green", RULES[1]: "amber", RULES[2]: "purple", RULES[3]: "red"})
md.conditional_formatting.add(f"Q{R0}:Q{RN}", FormulaRule(formula=[f'OR(LEFT(Q{R0},4)="Wait",LEFT(Q{R0},6)="Locked",LEFT(Q{R0},3)="Ask",LEFT(Q{R0},7)="Minimum")'], font=FD(AMBER, True, 9)))
md.conditional_formatting.add(f"R{R0}:R{RN}", CellIsRule(operator="equal", formula=['"Now"'], font=FD(GREEN, True)))
md.conditional_formatting.add(f"Q{R0}:Q{RN}", FormulaRule(formula=[f'OR(LEFT(Q{R0},5)="Worth",LEFT(Q{R0},2)="OK")'], font=FD(GREEN, False, 9)))
md.freeze_panes = f"C{R0}"
md[f"B{RN + 2}"] = "Rows marked (sample) are made up: overwrite or clear them. Up to 12 debts."; md[f"B{RN + 2}"].font = F(MUTED, False, 8, True)

# ======================= Calc sheets (one per order) =======================
FIRST = 12; LAST = FIRST + MONTHS - 1
rank_col = {"Calc Smart": "AA", "Calc Avalanche": "AB", "Calc Snowball": "AC", "Calc Cashflow": "AD"}
for sname in CALC:
    cs = wb.create_sheet(sname); cs.sheet_state = "hidden"
    rk = rank_col[sname]
    cs["A1"] = "row"; cs["A2"] = "rank"; cs["A3"] = "extra from"; cs["A4"] = "APR now"
    cs["A5"] = "debt-free"; cs["B5"] = f'=IFERROR(INDEX($B${FIRST}:$B${LAST},MATCH(1,INDEX(($AD${FIRST}:$AD${LAST}<=0.005)*1,0),0)),"-")'
    cs["A6"] = "months";    cs["B6"] = f'=IFERROR(MATCH(1,INDEX(($AD${FIRST}:$AD${LAST}<=0.005)*1,0),0),"-")'
    cs["A7"] = "interest";  cs["B7"] = f"=SUM($O${FIRST}:$Z${LAST})-(SUM($C$8:$N$8)-SUM($C${LAST}:$N${LAST}))"
    cs["A8"] = "fees";      cs["B8"] = "=SUM($C$1:$N$1)"
    cs["A9"] = "first";     cs["B9"] = '=IF(COUNT($C$10:$N$10)=0,"-",MIN($C$10:$N$10))'
    for j in range(N):
        c = L(3 + j); r = R0 + j
        cs[f"{c}2"] = f"=IF('My Debts'!${rk}${r}=\"\",\"\",'My Debts'!${rk}${r})"
        cs[f"{c}3"] = f"=IF('My Debts'!$U${r}=\"\",{FAR},'My Debts'!$U${r})"
        cs[f"{c}4"] = f"=N(+'My Debts'!$E${r})"
        cs[f"{c}5"] = f"=IF(N(+'My Debts'!$L${r})>0,'My Debts'!$L${r},{c}4)"
        cs[f"{c}6"] = f"=N(+'My Debts'!$K${r})"
        cs[f"{c}7"] = f"=N(+'My Debts'!$F${r})"
        cs[f"{c}8"] = f"=IF(AND(ISNUMBER('My Debts'!$D${r}),'My Debts'!$D${r}>0),'My Debts'!$D${r},0)"
        # payoff month for this debt, then early-payoff fee charged if cleared before the fee ends
        cs[f"{c}10"] = f'=IF({c}8<=0,"",IFERROR(INDEX($B${FIRST}:$B${LAST},MATCH(1,INDEX(({c}${FIRST}:{c}${LAST}<=0.005)*1,0),0)),""))'
        cs[f"{c}1"] = (f"=IF(AND('My Debts'!$H${r}=\"{RULES[1]}\",N(+'My Debts'!$I${r})>0,ISNUMBER({c}10)),"
                       f"IF(OR(N(+'My Debts'!$J${r})=0,{c}10<'My Debts'!$J${r}),'My Debts'!$I${r},0),0)")
    cs["B11"] = f"=EDATE({S['start']},-1)"
    for j in range(N): cs[f"{L(3 + j)}11"] = f"={L(3 + j)}8"
    cs["AD11"] = "=SUM(C11:N11)"
    for m in range(1, MONTHS + 1):
        r = FIRST + m - 1; p = r - 1
        cs[f"A{r}"] = m
        cs[f"B{r}"] = f"=EDATE({S['start']},{m - 1})"
        cs[f"AA{r}"] = f'=_xlfn.MINIFS($C$2:$N$2,C{p}:N{p},">0",$C$3:$N$3,"<="&B{r})'
        cs[f"AB{r}"] = f"=SUMPRODUCT((C{p}:N{p}>0)*$C$7:$N$7)"
        cs[f"AC{r}"] = f"=MAX(0,{S['budget']}-AB{r})+IF(A{r}={S['lumpm']},N({S['lump']}),0)"
        cs[f"AF{r}"] = f'=_xlfn.MINIFS($C$2:$N$2,C{p}:N{p},">0",$C$3:$N$3,"<="&B{r},$C$2:$N$2,">"&AA{r})'
        cs[f"AG{r}"] = f"=MAX(0,AC{r}-MAX(0,SUMPRODUCT(($C$2:$N$2=AA{r})*(AH{r}:AS{r}-$C$7:$N$7))))"
        for j in range(N):
            b, pc, dc = L(3 + j), L(15 + j), L(34 + j)
            cs[f"{dc}{r}"] = f"=IF({b}{p}<=0,0,ROUND({b}{p}*(1+IF(AND({b}$6>0,$B{r}<{b}$6),{b}$4,{b}$5)/12),2))"
            cs[f"{pc}{r}"] = f"=IF({b}{p}<=0,0,MIN({dc}{r},{b}$7+IF({b}$2=$AA{r},$AC{r},IF({b}$2=$AF{r},$AG{r},0))))"
            cs[f"{b}{r}"] = f"=IF({b}{p}<=0,0,MAX(0,{dc}{r}-{pc}{r}))"
        cs[f"AD{r}"] = f"=SUM(C{r}:N{r})"
        cs[f"AE{r}"] = f"=SUM(O{r}:Z{r})"

# ======================= Schedule (chosen order, visible) =======================
sc = sheet("Schedule", {"A": 3, "B": 12, "C": 14, "D": 14, "E": 10}, rows=8, cols=5 + N)
title(sc, "Month by month", "Your chosen order. Each column is one debt's balance at the end of the month.")
header(sc, 5, 2, ["Month", "Paid this month", "Total left", "Debts left"] + [f"Debt {k}" for k in range(1, N + 1)], ORANGE, 30)
for j in range(N):
    sc.cell(4, 6 + j, f"=IF('My Debts'!$B${R0 + j}=\"\",\"\",'My Debts'!$B${R0 + j})").font = F(MUTED, True, 8)
    sc.column_dimensions[L(6 + j)].width = 13
sc.cell(4, 6).alignment = Alignment(wrap_text=True)
pick = lambda col, r: "CHOOSE(" + S["idx"] + "," + ",".join(f"'{s}'!{col}{r}" for s in CALC) + ")"
for m in range(1, MONTHS + 1):
    r = 5 + m; cr = FIRST + m - 1; b = m % 2 == 0
    cell(sc, f"B{r}", f"=EDATE({S['start']},{m - 1})", MON, band=b)
    cell(sc, f"C{r}", "=" + pick("AE", cr), NUM, band=b, align="right")
    cell(sc, f"D{r}", "=" + pick("AD", cr), NUM, band=b, align="right")
    cell(sc, f"E{r}", "=" + "CHOOSE(" + S["idx"] + "," + ",".join(f"COUNTIF('{s}'!C{cr}:N{cr},\">0\")" for s in CALC) + ")", "0", band=b, align="center")
    for j in range(N):
        cell(sc, f"{L(6 + j)}{r}", "=" + pick(L(3 + j), cr), NUM, band=b, align="right", size=9)
sc.freeze_panes = "C6"
sc.conditional_formatting.add(f"D6:D{5 + MONTHS}", CellIsRule(operator="lessThanOrEqual", formula=["0.005"], font=FD(GREEN, True)))

# chart: total balance over time (first 10 years) on Plan
def rich(color, sz=900, bold=False):
    cp = CharacterProperties(sz=sz, b=bold, solidFill=color)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])
def dark(chart):
    chart.graphical_properties = GraphicalProperties(solidFill=PANEL, ln=LineProperties(noFill=True))
    chart.plot_area.graphicalProperties = GraphicalProperties(noFill=True, ln=LineProperties(noFill=True))
    for ax in (chart.x_axis, chart.y_axis):
        ax.delete = False; ax.txPr = rich(MUTED, 800); ax.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=LINE))
    chart.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=LINE))
lc = LineChart(); lc.legend = None
lc.add_data(Reference(sc, min_col=4, min_row=5, max_row=5 + 120), titles_from_data=True)
lc.set_categories(Reference(sc, min_col=2, min_row=6, max_row=5 + 120))
lc.series[0].graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=ORANGE, w=28000)); lc.series[0].smooth = False
lc.x_axis.number_format = "mmm yy"; lc.y_axis.number_format = "#,##0"; dark(lc)
lc.title = None; lc.width, lc.height = 15, 7
pl["H13"] = "TOTAL LEFT, FIRST 10 YEARS"; pl["H13"].font = F(ORANGE, True, 9)
pl.add_chart(lc, "H14")
bc = BarChart(); bc.type = "col"; bc.legend = None; bc.gapWidth = 60
bc.add_data(Reference(pl, min_col=3, max_col=6, min_row=19), from_rows=True, titles_from_data=False)
bc.set_categories(Reference(pl, min_col=3, max_col=6, min_row=14))
bc.series[0].graphicalProperties = GraphicalProperties(solidFill=TEAL, ln=LineProperties(noFill=True))
bc.y_axis.number_format = "#,##0"; dark(bc); bc.width, bc.height = 15, 6.5
pl["H29"] = "INTEREST + FEES BY ORDER"; pl["H29"].font = F(TEAL, True, 9)
pl.add_chart(bc, "H46")

# ======================= Start Here =======================
hw = sheet("Start Here", {"A": 3, "B": 112}, rows=30, cols=3)
title(hw, "Start here", "10 minutes to set up. Works with any currency.")
steps = ["1. My Debts: one row per debt. Fill the blue columns. Only Name, Balance, APR and Minimum payment are needed.",
         "   • Optional columns can stay blank: blank = no fee, no lock, no promo, not behind. The maths still works.",
         "   • Early payoff rule: if your contract charges a fee for paying early, pick 'Allowed with a fee' and enter the fee and when it ends.",
         "     The sheet compares the fee with the interest you'd save. If the fee costs more, extra payments wait until the fee ends.",
         "   • Locked until a date: no extra goes in until that date (minimums still get paid).",
         "   • Not sure: minimums only, until you ask your lender and change the rule.",
         "   • 0% / promo deals: enter the promo end date and the APR after it. Deferred interest = Yes if back-interest is added when you don't clear it in time.",
         "2. Plan: enter how much you can put toward all debts each month, the start month, and any one-time extra.",
         "3. Plan: pick a payoff order and compare all four. Smart is CheckMaybe's default:",
         "   behind-on-payments first → deferred-interest deadlines → highest APR (or, if minimums take 35%+ of income, the debt that frees up the most monthly money).",
         "4. Schedule: month by month, what's left on each debt.",
         "",
         "How the extra money moves: every month all minimums get paid; whatever's left goes to the #1 debt that is allowed to take extra.",
         "When a debt is cleared, its minimum rolls into the next one. Interest is APR ÷ 12 each month.",
         "",
         "Estimates only. Not financial advice. Contracts differ by country and lender: your lender's figures are the final word.",
         "Example rows are made up. Replace or clear them."]
for i, s in enumerate(steps):
    hw[f"B{5 + i}"] = s; hw[f"B{5 + i}"].font = F(TXT if s[:1].isdigit() else MUTED, s[:1].isdigit(), 10)
    hw[f"B{5 + i}"].alignment = Alignment(wrap_text=True, vertical="top")
wb.move_sheet("Start Here", -(len(wb.sheetnames) - 1))
order = ["Start Here", "Plan", "My Debts", "Schedule"] + CALC
wb._sheets = [wb[n] for n in order]
wb.active = 1
for ws in wb.worksheets: ws.sheet_properties.tabColor = {"Plan": ORANGE, "Start Here": TEAL, "My Debts": BLUE}.get(ws.title, "3A3F58")
wb.save("Debt-Payoff-Planner.xlsx"); import sys, os; sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "build")); import excel_compat; excel_compat.fix("Debt-Payoff-Planner.xlsx"); print("saved", wb.sheetnames)
